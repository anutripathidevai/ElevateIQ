"""Shared mechanical checks. Never imports supplied assessment implementations."""
from __future__ import annotations

import ast
from collections import Counter
import csv
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import platform
import re
import sys
import zipfile

import openpyxl
from file_io import retry_io, write_json

MODELS = ("reference", "a", "b", "c")
HEADERS = (
    "Task ID", "Category", "Prompt", "Expert Rubric", "Hidden Tests",
    "Model A Tests Passed", "Model A Rubric Score", "Model B Tests Passed",
    "Model B Rubric Score", "Model C Tests Passed", "Model C Rubric Score",
)


CRITERIA = ("description", "importance", "points", "type", "rationale", "source")


def digest(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest() if hasattr(hashlib, "file_digest") else _digest_stream(stream)


def _digest_stream(stream):
    value = hashlib.sha256()
    for chunk in iter(lambda: stream.read(1024 * 1024), b""):
        value.update(chunk)
    return value.hexdigest()


def safe_path(root, relative):
    path = (root / relative).resolve()
    if not path.is_relative_to(root.resolve()):
        raise ValueError(f"Path escapes root: {relative}")
    return path


def save_json(path, value):
    write_json(path, value)


def save_csv(path, rows, fields):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="raise")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: json.dumps(value, ensure_ascii=True) if isinstance(value, (dict, list)) else value
                             for key, value in row.items()})
    retry_io(lambda: temporary.replace(path))


def load_json(path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def observation(task, location, check, status, message):
    content = json.dumps([task, location, check, message], sort_keys=True)
    return dict(evidence_id=hashlib.sha256(content.encode()).hexdigest()[:16],
                task_id=task, location=location, check=check, status=status, message=message)


def envelope(ctx, records):
    return dict(schema_version=1, run_id=ctx["run_id"], source_snapshot=ctx.get("snapshot"),
                root=str(ctx["root"]), generated_utc=datetime.now(timezone.utc).isoformat(),
                command=sys.argv, python=platform.python_version(), records=records)


def files_in(root):
    files, observations = [], []
    def walk(directory):
        try:
            children = sorted(directory.iterdir(), key=lambda p: p.name)
        except OSError as error:
            observations.append(observation("", str(directory), "directory_read", "blocked", str(error)))
            return
        for path in children:
            relative = path.relative_to(root).as_posix()
            if path.is_symlink() or not path.resolve().is_relative_to(root):
                observations.append(observation("", relative, "link", "warning", "Link skipped; target not read."))
            elif path.is_dir():
                walk(path)
            elif path.is_file():
                try:
                    files.append(dict(path=relative, bytes=path.stat().st_size, sha256=digest(path), error=""))
                except OSError as error:
                    files.append(dict(path=relative, bytes=None, sha256=None, error=str(error)))
    walk(root)
    return files, observations


def source_rows(root):
    """Keep row/cell identity and raw formula text. Never calculate workbook formulas."""
    path = root / "EvalTasks.xlsx"
    rows, issues = [], []
    if not path.is_file():
        return rows, [observation("", "EvalTasks.xlsx", "workbook_access", "blocked", "Workbook absent; directory inventory remains available.")]
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as error:
        return rows, [observation("", "EvalTasks.xlsx", "workbook_access", "blocked", str(error))]
    try:
        for sheet in workbook:
            iterator = sheet.iter_rows()
            first = next(iterator, ())
            names = [cell.value for cell in first]
            mapping = {name: index for index, name in enumerate(names) if name is not None}
            for name in HEADERS:
                count = names.count(name)
                if count != 1:
                    issues.append(observation("", f"EvalTasks.xlsx::{sheet.title}!1", "header", "failure",
                                              f"Header {name!r}: expected once, found {count}."))
            if "Task ID" not in mapping:
                continue
            for number, cells in enumerate(iterator, 2):
                if all(cell.value is None for cell in cells):
                    continue
                values = {name: cells[index].value if index < len(cells) else None for name, index in mapping.items()}
                for name, value in list(values.items()):
                    if value is not None and type(value) not in (str, bool, int, float):
                        issues.append(observation("", f"EvalTasks.xlsx::{sheet.title}!{number}", "cell_type", "warning",
                                                  f"{name}: unsupported cell type {type(value).__name__}; text representation retained."))
                        values[name] = repr(value)
                    elif type(value) is float and not math.isfinite(value):
                        values[name] = repr(value)
                raw_id = values.get("Task ID")
                task_id = str(raw_id).strip() if raw_id is not None else ""
                if not re.fullmatch(r"[0-9]+", task_id):
                    issues.append(observation(task_id, f"EvalTasks.xlsx::{sheet.title}!{number}", "task_id", "failure",
                                              f"Non-numeric or blank task identifier {raw_id!r}; row retained."))
                locations = {name: f"EvalTasks.xlsx::{sheet.title}!{openpyxl.utils.get_column_letter(index+1)}{number}"
                             for name, index in mapping.items()}
                rows.append(dict(task_id=task_id, sheet=sheet.title, row=number, values=values, locations=locations))
    finally:
        workbook.close()
    return rows, issues


def task_layout(task_id):
    return dict(prompt=f"Data/Task_{task_id}/prompt.md", starter=f"Data/Task_{task_id}/solution.py",
                tests=f"Data/ExpertSolution_{task_id}/tests",
                implementations={model: f"Data/ExpertSolution_{task_id}/reference/solution.py" if model == "reference"
                                 else f"Data/ModelOutputs_{task_id}/model_{model}/solution.py" for model in MODELS})


def is_generated_cache(relative: str) -> bool:
    parts = Path(relative).parts
    return "__pycache__" in parts or ".pytest_cache" in parts or Path(relative).suffix == ".pyc"


def inventory(ctx):
    root = ctx["root"]
    manifest, issues = files_in(root)
    for name in ("README_TakeHomeAssessment.pdf", "ExpertGuidelines.pdf"):
        if not (root / name).is_file():
            issues.append(observation("", name, "instructions_presence", "warning", "Source instructions not found under selected assessment root."))
    rows, workbook_issues = source_rows(root)
    issues.extend(workbook_issues)
    directories = {}
    data = root / "Data"
    if data.is_dir():
        for path in sorted(data.iterdir()):
            if path.is_dir():
                match = re.fullmatch(r"(Task|ExpertSolution|ModelOutputs)_([0-9]+)", path.name)
                if match:
                    directories.setdefault(match[2], {})[match[1]] = path.relative_to(root).as_posix()
                else:
                    issues.append(observation("", path.relative_to(root).as_posix(), "directory_mapping", "warning", "Directory outside documented naming pattern."))
    counts = Counter(row["task_id"] for row in rows)
    for task_id, count in sorted(counts.items()):
        if count > 1:
            issues.append(observation(task_id, "EvalTasks.xlsx", "duplicate_id", "failure", f"{count} workbook rows share this ID; no row selected silently."))
    ids = sorted(set(directories) | {i for i in counts if re.fullmatch(r"[0-9]+", i)})
    tasks = []
    known_files = set()
    for task_id in ids:
        layout = task_layout(task_id)
        task_rows = [row for row in rows if row["task_id"] == task_id]
        prompt_value = task_rows[0]["values"].get("Prompt", "") if len(task_rows) == 1 else ""
        heading = re.search(r"(?m)^#\s+([^\r\n]+)", prompt_value) if isinstance(prompt_value, str) else None
        tasks.append(dict(task_id=task_id, task_name=heading[1] if heading else None,
                          task_name_source=task_rows[0]["locations"].get("Prompt") if heading else None,
                          layout=layout, workbook_rows=task_rows, directories=directories.get(task_id, {})))
        expected = [layout["prompt"], layout["starter"], *layout["implementations"].values(), layout["tests"] + "/test_hidden.py"]
        known_files.update(expected)
        for relative in expected:
            if not safe_path(root, relative).is_file():
                conditional = relative == layout["starter"]
                issues.append(observation(task_id, relative, "file_presence", "warning" if conditional else "failure",
                                          "Starter file absent; applicability conditional." if conditional else "Documented baseline path absent."))
        if len(task_rows) != 1 or len(directories.get(task_id, {})) != 3:
            issues.append(observation(task_id, "Data", "mapping", "warning", "Expected one workbook row and three directory groups; inspect recorded mappings."))
    if ctx["config"].get("expected_tasks") is not None and (
        len(ids) != ctx["config"]["expected_tasks"] or len(rows) != ctx["config"]["expected_tasks"]
    ):
        issues.append(observation("", "Data;EvalTasks.xlsx", "expected_count", "warning",
                                  f"Expected {ctx['config']['expected_tasks']}; observed {len(ids)} mapped IDs and {len(rows)} workbook rows."))
    for entry in manifest:
        if entry["error"]:
            issues.append(observation("", entry["path"], "file_read", "blocked", entry["error"]))
        if entry["path"].startswith("Data/") and entry["path"] not in known_files:
            cache = is_generated_cache(entry["path"])
            issues.append(observation(
                "", entry["path"], "generated_cache" if cache else "additional_file",
                "info" if cache else "warning",
                "Generated Python/pytest cache inventoried but not copied into execution snapshots; original left intact."
                if cache else "Additional source file retained; no purpose or preferred version inferred.",
            ))
    snapshot = hashlib.sha256(json.dumps(manifest, sort_keys=True).encode()).hexdigest()
    ctx["snapshot"] = snapshot
    result = envelope(ctx, dict(tasks=tasks, workbook_rows=rows, manifest=manifest, observations=issues))
    save_json(ctx["out"] / "inventory.json", result)
    return issues


def read_inventory(ctx):
    value = load_json(ctx["out"] / "inventory.json")
    if Path(value["root"]).resolve() != ctx["root"]:
        raise ValueError("Inventory belongs to a different root. Select the matching --output.")
    ctx["snapshot"] = value["source_snapshot"]
    return value["records"]


def integrity(ctx, inventory_value):
    observed, issues = files_in(ctx["root"])
    previous = {entry["path"]: entry for entry in inventory_value["manifest"]}
    current = {entry["path"]: entry for entry in observed}
    for path in sorted(set(previous) | set(current)):
        if path not in current or path not in previous or current[path]["sha256"] != previous[path]["sha256"] or current[path]["error"]:
            issues.append(observation("", path, "source_integrity", "failure", "Source added, removed, unreadable, or changed since original manifest."))
    return issues


def parse_rubric(text):
    def reject(value):
        raise ValueError(f"Non-JSON constant: {value}")
    if not isinstance(text, str):
        return None, [("rubric_json", "failure", "Rubric cell must contain JSON text.")]
    try:
        value = json.loads(text, parse_constant=reject)
    except (ValueError, TypeError) as error:
        return None, [("rubric_json", "failure", str(error))]
    if not isinstance(value, list):
        return None, [("rubric_array", "failure", "JSON root is not an array.")]
    errors = []
    for index, item in enumerate(value):
        if not isinstance(item, dict):
            errors.append(("criterion_object", "failure", f"Criterion {index} is not an object."))
            continue
        for key in CRITERIA:
            if key not in item:
                errors.append(("required_field", "failure", f"Criterion {index}: missing {key}."))
        for key in ("description", "importance", "rationale", "source"):
            if key in item and not isinstance(item[key], str):
                errors.append(("field_type", "warning", f"Criterion {index}.{key}: expected text based on documented example; no formal schema supplied."))
        if "points" in item and not numeric(item["points"]):
            errors.append(("points_type", "failure", f"Criterion {index}.points must be a finite number for arithmetic."))
        if "type" in item and item["type"] not in ("tests", "code_review"):
            errors.append(("criterion_type", "failure", f"Criterion {index}.type is not tests or code_review."))
    return value, errors


def numeric(value):
    return type(value) in (int, float) and math.isfinite(value)


def artifacts(ctx):
    inv = read_inventory(ctx)
    issues = integrity(ctx, inv)
    # Re-read workbook rather than relying on a previously extracted cell value.
    current_rows, workbook_issues = source_rows(ctx["root"])
    issues.extend(workbook_issues)
    if current_rows != inv["workbook_rows"]:
        issues.append(observation("", "EvalTasks.xlsx", "workbook_snapshot", "failure", "Workbook rows differ from inventory; refresh a new inventory before baseline execution."))
    manifest = {entry["path"]: entry for entry in inv["manifest"]}
    for row in current_rows:
        task_id = row["task_id"]
        rubric, errors = parse_rubric(row["values"].get("Expert Rubric"))
        location = row["locations"].get("Expert Rubric", "EvalTasks.xlsx")
        issues.extend(observation(task_id, location, check, status, message) for check, status, message in errors)
        for field in HEADERS[4:]:
            value = row["values"].get(field)
            if not numeric(value):
                issues.append(observation(task_id, row["locations"].get(field, "EvalTasks.xlsx"), "recorded_numeric_value", "warning",
                                          f"{field}: missing, formula, or non-finite/non-numeric value; not substituted."))
            elif field == "Hidden Tests" or "Tests Passed" in field:
                if value < 0 or int(value) != value:
                    issues.append(observation(task_id, row["locations"][field], "count_value", "failure", "Test count must be a nonnegative integer."))
        hidden = row["values"].get("Hidden Tests")
        for model in ("A", "B", "C"):
            field = f"Model {model} Tests Passed"
            passed = row["values"].get(field)
            if numeric(hidden) and numeric(passed) and passed > hidden:
                issues.append(observation(task_id, row["locations"][field], "count_relation", "failure", "Recorded tests passed exceeds recorded hidden tests."))
        if rubric:
            for index, criterion in enumerate(rubric):
                source = criterion.get("source") if isinstance(criterion, dict) else None
                if not isinstance(source, str):
                    continue
                # Literal path tokens only; descriptive references remain unresolved.
                for token in re.findall(r"[\w./\\-]+\.(?:py|md)\b", source):
                    token = token.replace("\\", "/")
                    possible = [token, f"Data/ExpertSolution_{task_id}/{token}",
                                f"Data/Task_{task_id}/{token}"]
                    if token.startswith("task/"):
                        possible.append(f"Data/Task_{task_id}/{token[5:]}")
                    if not any(path in manifest for path in possible):
                        issues.append(observation(task_id, f"{location}::criterion[{index}].source", "source_reference", "warning",
                                                  f"Literal reference {token!r} not resolved against inventory; no substitute created."))
    for task in inv["tasks"]:
        task_id, layout = task["task_id"], task["layout"]
        prompt = safe_path(ctx["root"], layout["prompt"])
        if prompt.is_file() and len(task["workbook_rows"]) == 1:
            try:
                text = prompt.read_text(encoding="utf-8-sig")
            except (OSError, UnicodeError) as error:
                issues.append(observation(task_id, layout["prompt"], "prompt_read", "failure", str(error)))
            else:
                source = task["workbook_rows"][0]["values"].get("Prompt")
                if isinstance(source, str) and text.replace("\r\n", "\n") != source.replace("\r\n", "\n"):
                    cosmetic = text.strip() == source.strip()
                    issues.append(observation(
                        task_id, layout["prompt"],
                        "prompt_outer_whitespace" if cosmetic else "prompt_equality",
                        "info" if cosmetic else "warning",
                        "Prompt texts agree after trimming outer whitespace." if cosmetic else
                        "File and workbook prompt text differ beyond outer whitespace; authority requires manual review.",
                    ))
        test_paths = [path for path in manifest if path.startswith(layout["tests"] + "/") and path.endswith(".py")
                      and (Path(path).name.startswith("test_") or Path(path).name.endswith("_test.py"))]
        if len(test_paths) > 1:
            issues.append(observation(task_id, layout["tests"], "test_discovery", "warning", f"Multiple pytest-named test files: {test_paths}; all retained in baseline."))
    dependencies = []
    duplicate_groups = {}
    for relative, entry in manifest.items():
        if relative.startswith("Data/") and entry["sha256"] and entry["bytes"]:
            duplicate_groups.setdefault(entry["sha256"], []).append(relative)
    for paths in duplicate_groups.values():
        if len(paths) > 1:
            issues.append(observation(
                "", ";".join(sorted(paths)), "identical_artifacts", "info",
                "Byte-identical supplied files; duplication alone does not establish a dataset defect.",
            ))
    for relative in sorted(manifest):
        if not relative.endswith(".py"):
            continue
        path = safe_path(ctx["root"], relative)
        match = re.search(r"(?:Task|ExpertSolution|ModelOutputs)_([0-9]+)", relative)
        task_id = match[1] if match else ""
        try:
            tree = ast.parse(path.read_bytes(), filename=relative)
        except (OSError, SyntaxError, ValueError) as error:
            issues.append(observation(task_id, relative, "python_parse", "failure", str(error)))
            continue
        for node in ast.walk(tree):
            modules = [alias.name for alias in node.names] if isinstance(node, ast.Import) else [node.module or ""] if isinstance(node, ast.ImportFrom) else []
            for module in modules:
                top = module.split(".")[0]
                resolved = top in sys.stdlib_module_names or top in ("pytest", "solution") or (path.parent / f"{top}.py").is_file() or (path.parent / top / "__init__.py").is_file()
                item = dict(task_id=task_id, path=relative, line=node.lineno, module=module, level=getattr(node, "level", 0),
                            status="known_or_adjacent" if resolved else "unresolved_static")
                dependencies.append(item)
                if not resolved:
                    issues.append(observation(task_id, f"{relative}:{node.lineno}", "import_resolution", "warning", f"Import {module!r} not resolved statically; no import attempted."))
    save_json(ctx["out"] / "artifact_validation.json", envelope(ctx, dict(observations=issues, imports=dependencies)))
    return issues


def baseline_data(ctx, filename="baseline_results.json"):
    value = load_json(ctx["out"] / filename)
    if value["source_snapshot"] != ctx["snapshot"] or Path(value["root"]).resolve() != ctx["root"]:
        raise ValueError("Execution results do not match selected source snapshot.")
    return value["records"]


def reconcile(ctx):
    inv = read_inventory(ctx)
    results = baseline_data(ctx)
    current_rows, issues = source_rows(ctx["root"])
    issues.extend(integrity(ctx, inv))
    by_id = {}
    for row in current_rows:
        by_id.setdefault(row["task_id"], []).append(row)
    comparisons = []
    for run in results:
        rows = by_id.get(run["task_id"], [])
        row = rows[0] if len(rows) == 1 else None
        model = run["model"]
        specs = [("Hidden Tests", "collected")]
        if model != "reference":
            specs.append((f"Model {model.upper()} Tests Passed", "passed"))
        for field, measured in specs:
            recorded = row["values"].get(field) if row else None
            counts = run.get("counts") or {}
            comparable = run["status"] in ("passed", "test_failed") and numeric(recorded) and not issues
            actual = counts.get(measured)
            status = "match" if comparable and recorded == actual else "difference" if comparable and actual is not None else "incomparable"
            comparisons.append(dict(task_id=run["task_id"], model=model, metric=field,
                                    workbook_location=row["locations"].get(field, "") if row else "",
                                    workbook_value=recorded, observed_value=actual,
                                    execution_status=run["status"], status=status, evidence=run["log"],
                                    note="No causal interpretation. Hidden Tests compared to collected pytest items, including parametrization."))
    save_json(ctx["out"] / "result_reconciliation.json", envelope(ctx, dict(comparisons=comparisons, observations=issues)))
    return issues + [observation(row["task_id"], row["evidence"], "result_comparison",
                                 "failure" if row["status"] == "difference" else "blocked",
                                 f"{row['metric']}: {row['status']}")
                     for row in comparisons if row["status"] != "match"]


def scoring(ctx):
    inv = read_inventory(ctx)
    issues = integrity(ctx, inv)
    rows, workbook_issues = source_rows(ctx["root"])
    issues.extend(workbook_issues)
    awards, rules = {}, {}
    if ctx["config"].get("criterion_awards"):
        with Path(ctx["config"]["criterion_awards"]).open(newline="", encoding="utf-8-sig") as stream:
            for item in csv.DictReader(stream):
                key = (item["task_id"], item["model"], int(item["criterion_index"]))
                if key[1] not in ("a", "b", "c") or key[2] < 0:
                    raise ValueError(f"Invalid criterion award model/index: {key}")
                if key in awards:
                    raise ValueError(f"Duplicate criterion award: {key}")
                number = float(item["awarded_points"])
                if not numeric(number):
                    raise ValueError(f"Non-finite criterion award: {key}")
                awards[key] = number
    if ctx["config"].get("scoring_rules"):
        rules = load_json(Path(ctx["config"]["scoring_rules"]))
        if not isinstance(rules, dict):
            raise ValueError("scoring_rules must be an object keyed by task ID.")
    observed = {}
    if (ctx["out"] / "baseline_results.json").is_file():
        observed = {(row["task_id"], row["model"]): row for row in baseline_data(ctx)}
    for key in awards:
        matching = [row for row in rows if row["task_id"] == key[0]]
        if len(matching) != 1:
            raise ValueError(f"Criterion award does not map to exactly one workbook row: {key}")
        criteria, errors = parse_rubric(matching[0]["values"].get("Expert Rubric"))
        if criteria is not None and key[2] >= len(criteria):
            raise ValueError(f"Criterion award index exceeds supplied rubric: {key}")
    table, schema_errors = [], []
    for row in rows:
        rubric, errors = parse_rubric(row["values"].get("Expert Rubric"))
        schema_errors.extend(observation(row["task_id"], row["locations"].get("Expert Rubric", "EvalTasks.xlsx"), check, status, message)
                             for check, status, message in errors)
        for model in ("a", "b", "c"):
            task_id = row["task_id"]
            rule = rules.get(task_id, {})
            if not isinstance(rule, dict):
                raise ValueError(f"Scoring rule for {task_id} must be an object.")
            recorded = row["values"].get(f"Model {model.upper()} Rubric Score")
            hidden = row["values"].get("Hidden Tests")
            passed = row["values"].get(f"Model {model.upper()} Tests Passed")
            simultaneous = numeric(passed) and numeric(hidden) and passed < hidden
            observed_run = observed.get((task_id, model), {})
            observed_counts = observed_run.get("counts") or {}
            observed_failures = None
            if observed_run.get("status") in ("passed", "test_failed"):
                observed_failures = observed_counts.get("failed")
            status, total, note = "unavailable", None, "Criterion awards and explicit supported scoring rule required; no formula inferred."
            if rubric is not None and not errors and rule.get("formula") == "sum_awards":
                if not rule.get("source"):
                    raise ValueError(f"Scoring rule for {task_id} requires a source reference.")
                selected = [awards.get((task_id, model, i)) for i in range(len(rubric))]
                if all(value is not None for value in selected):
                    total = sum(selected)
                    status = "match" if numeric(recorded) and recorded == total else "difference" if numeric(recorded) else "incomparable"
                    note = f"Summed supplied criterion awards; rule source: {rule['source']}. No code-review scores generated."
                    ranges = rule.get("award_ranges", {})
                    if not isinstance(ranges, dict):
                        raise ValueError(f"award_ranges for {task_id} must be an object.")
                    for index, value in enumerate(selected):
                        bounds = ranges.get(str(index))
                        if bounds is not None and (not isinstance(bounds, list) or len(bounds) != 2 or
                                                   not all(numeric(number) for number in bounds) or bounds[0] > bounds[1]):
                            raise ValueError(f"Invalid award range for {task_id}, criterion {index}: {bounds}")
                        if bounds is not None and not bounds[0] <= value <= bounds[1]:
                            schema_errors.append(observation(task_id, f"criterion_awards::{model}[{index}]", "explicit_award_range", "failure", f"Award {value} outside supplied range {bounds}."))
            threshold = ctx["config"].get("high_score_threshold")
            table.append(dict(task_id=task_id, model=model, workbook_location=row["locations"].get(f"Model {model.upper()} Rubric Score"),
                              recorded_score=recorded, recorded_tests_passed=passed, recorded_hidden_tests=hidden,
                              recorded_nonpasses_present=simultaneous,
                              observed_execution_status=observed_run.get("status"),
                              observed_failed_tests=observed_failures,
                              high_score_threshold=threshold,
                              threshold_with_recorded_nonpasses=bool(threshold is not None and numeric(recorded) and recorded >= threshold and simultaneous),
                              reconstructed_score=total, reconstruction_status=status, note=note))
    save_json(ctx["out"] / "scoring_validation.json", envelope(
        ctx, dict(scores=table, observations=issues + schema_errors)))
    return issues + schema_errors + [observation(row["task_id"], row["workbook_location"], "score_reconstruction", "warning" if row["reconstruction_status"] == "unavailable" else "failure", row["note"])
                                    for row in table if row["reconstruction_status"] in ("unavailable", "difference")]
